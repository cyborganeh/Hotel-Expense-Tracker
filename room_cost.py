"""Daily room-cost input, MTD calculations, and local persistence."""
import json
import os
import re
import hashlib
import sqlite3
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
import pandas as pd


DAILY_COLUMNS = ["Date", "Occupancy", "Room_Nights", "Water", "Electricity", "Linen_Laundry"]
REQUIRED_DAILY_COLUMNS = ["Date", "Room_Nights", "Water", "Electricity", "Linen_Laundry"]
EXPENSE_COLUMNS = ["Water", "Electricity", "Linen_Laundry"]


def _month_key(month: str) -> str:
    value = str(month).strip()
    if not re.fullmatch(r"\d{4}-\d{2}", value):
        raise ValueError("Month must use YYYY-MM format.")
    year, month_number = (int(part) for part in value.split("-"))
    if not 1 <= month_number <= 12:
        raise ValueError("Month must use YYYY-MM format.")
    return value


def _normalise_daily_frame(
    frame: pd.DataFrame,
    month: str,
    *,
    allow_missing_expenses: bool = False,
    allow_missing_room_nights: bool = False,
) -> pd.DataFrame:
    month = _month_key(month)
    if frame.empty:
        return pd.DataFrame(columns=DAILY_COLUMNS)

    aliases = {
        "date": "Date",
        "occupancy": "Occupancy",
        "occupancypercent": "Occupancy",
        "occ": "Occupancy",
        "occpercent": "Occupancy",
        "roomnights": "Room_Nights",
        "roomsold": "Room_Nights",
        "expocc": "Room_Nights",
        "water": "Water",
        "watercost": "Water",
        "electricity": "Electricity",
        "electricitycost": "Electricity",
        "linen": "Linen_Laundry",
        "laundry": "Linen_Laundry",
        "linenlaundry": "Linen_Laundry",
        "linenlaundrycost": "Linen_Laundry",
    }
    renamed = {}
    for column in frame.columns:
        key = re.sub(r"[^a-z0-9]", "", str(column).strip().lower())
        if key in aliases:
            canonical = aliases[key]
            if canonical in renamed.values():
                raise ValueError(f"Excel contains more than one column for {canonical.replace('_', ' ')}.")
            renamed[column] = canonical
    normalized = frame.rename(columns=renamed)

    if "Occupancy" not in normalized.columns:
        normalized["Occupancy"] = pd.NA
    missing_columns = [column for column in REQUIRED_DAILY_COLUMNS if column not in normalized.columns]
    if missing_columns:
        display = ", ".join(column.replace("_", " ") for column in missing_columns)
        raise ValueError(f"Excel is missing required columns: {display}.")

    data = normalized[DAILY_COLUMNS].copy()
    parsed_dates = pd.to_datetime(data["Date"], errors="coerce", format="mixed")
    invalid_dates = parsed_dates.isna()
    if invalid_dates.any():
        rows = ", ".join(str(index + 2) for index in data.index[invalid_dates])
        raise ValueError(f"Invalid or missing Date value on Excel row(s): {rows}.")
    data["Date"] = parsed_dates.dt.strftime("%Y-%m-%d")

    wrong_month = data["Date"].str[:7] != month
    if wrong_month.any():
        dates = ", ".join(data.loc[wrong_month, "Date"].tolist())
        raise ValueError(f"Date(s) outside selected month {month}: {dates}.")

    duplicate_dates = data.loc[data["Date"].duplicated(keep=False), "Date"].unique().tolist()
    if duplicate_dates:
        raise ValueError(f"Excel has duplicate date(s): {', '.join(duplicate_dates)}.")

    for column in ["Room_Nights", *EXPENSE_COLUMNS]:
        if column == "Room_Nights" and allow_missing_room_nights:
            values = pd.to_numeric(data[column], errors="coerce")
            if (values.dropna() < 0).any() or (values.dropna() % 1 != 0).any():
                raise ValueError("Room Nights must be a positive whole number when provided.")
            if (values.dropna() == 0).any():
                raise ValueError("Room Nights must be positive; zero is not a valid value.")
            data[column] = values.astype("Int64")
            continue
        if column in EXPENSE_COLUMNS and allow_missing_expenses:
            values = pd.to_numeric(data[column], errors="coerce")
            if (values.dropna() < 0).any():
                raise ValueError(f"{column.replace('_', ' ')} cannot be negative.")
            data[column] = values.astype(float)
            continue
        values = pd.to_numeric(data[column], errors="coerce")
        if values.isna().any():
            rows = ", ".join(str(index + 2) for index in data.index[values.isna()])
            raise ValueError(f"Missing or non-numeric {column.replace('_', ' ')} on Excel row(s): {rows}.")
        if (values < 0).any():
            raise ValueError(f"{column.replace('_', ' ')} cannot be negative.")
        data[column] = values.astype(float)

    occupancy = pd.to_numeric(data["Occupancy"], errors="coerce")
    if (occupancy.dropna() < 0).any() or (occupancy.dropna() > 100).any():
        raise ValueError("Occupancy must be between 0 and 100 percent.")
    data["Occupancy"] = occupancy.astype(float)

    invalid_room_nights = (data["Room_Nights"] <= 0) | (data["Room_Nights"] % 1 != 0)
    if invalid_room_nights.any():
        raise ValueError("Room Nights must be a positive whole number for every date.")

    data["Room_Nights"] = data["Room_Nights"].astype("Int64" if allow_missing_room_nights else int)
    return data.sort_values("Date").reset_index(drop=True)


def parse_daily_room_cost_excel(source, month: str) -> pd.DataFrame:
    """Read and validate daily room-cost rows from an uploaded Excel workbook."""
    if hasattr(source, "seek"):
        source.seek(0)
    frame = pd.read_excel(source, engine="openpyxl")
    if hasattr(source, "seek"):
        source.seek(0)
    return _normalise_daily_frame(frame, month)


def parse_night_audit_report(source, month: Optional[str] = None) -> pd.DataFrame:
    """Extract date, occupancy, and expected rooms occupied from a Night Audit report."""
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    try:
        for worksheet in workbook.worksheets:
            header_row = None
            columns = {}
            for row_idx, values in enumerate(worksheet.iter_rows(values_only=True), 1):
                labels = {
                    re.sub(r"[^a-z0-9]", "", str(value).strip().lower()): col_idx
                    for col_idx, value in enumerate(values)
                    if value is not None
                }
                date_col = labels.get("date")
                room_sold_col = labels.get("expocc") or labels.get("roomsold")
                occupancy_col = labels.get("occ") or labels.get("occpercent")
                if date_col is not None and room_sold_col is not None and occupancy_col is not None:
                    header_row = row_idx
                    columns = {
                        "Date": date_col,
                        "Room_Nights": room_sold_col,
                        "Occupancy": occupancy_col,
                    }
                    break
            if header_row is None:
                continue

            records = []
            for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                date_value = values[columns["Date"]] if columns["Date"] < len(values) else None
                if isinstance(date_value, datetime):
                    parsed_date = pd.Timestamp(date_value).strftime("%Y-%m-%d")
                else:
                    parsed = pd.to_datetime(date_value, errors="coerce", format="mixed")
                    if pd.isna(parsed):
                        continue
                    parsed_date = parsed.strftime("%Y-%m-%d")

                room_sold = values[columns["Room_Nights"]] if columns["Room_Nights"] < len(values) else None
                occupancy = values[columns["Occupancy"]] if columns["Occupancy"] < len(values) else None
                if room_sold is None or occupancy is None:
                    continue
                records.append({
                    "Date": parsed_date,
                    "Occupancy": occupancy,
                    "Room_Nights": room_sold,
                    "Water": pd.NA,
                    "Electricity": pd.NA,
                    "Linen_Laundry": pd.NA,
                })
            if records:
                detected_month = month or records[0]["Date"][:7]
                return _normalise_daily_frame(
                    pd.DataFrame(records), detected_month, allow_missing_expenses=True
                )
        raise ValueError("Could not find DATE, EXP OCC/ROOM SOLD, and OCC (%) columns in the workbook.")
    finally:
        workbook.close()


def parse_energy_cost_workbook(source, month: str) -> pd.DataFrame:
    """Extract daily electricity and water amounts from an Energy Cost workbook."""
    month = _month_key(month)
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    try:
        for worksheet in workbook.worksheets:
            header_row = None
            columns = {}
            for row_idx, values in enumerate(worksheet.iter_rows(values_only=True), 1):
                labels = {
                    re.sub(r"[^a-z0-9]", "", str(value).strip().lower()): col_idx
                    for col_idx, value in enumerate(values)
                    if value is not None
                }
                date_col = labels.get("tanggal")
                if date_col is None:
                    date_col = labels.get("date")
                electricity_col = next(
                    (idx for key, idx in labels.items() if "listrik" in key or "electricity" in key),
                    None,
                )
                water_col = next(
                    (idx for key, idx in labels.items() if "air" in key or "water" in key),
                    None,
                )
                if date_col is not None and electricity_col is not None and water_col is not None:
                    header_row = row_idx
                    columns = {"Date": date_col, "Electricity": electricity_col, "Water": water_col}
                    break
            if header_row is None:
                continue

            year, month_number = (int(part) for part in month.split("-"))
            records = []
            for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                day_value = values[columns["Date"]] if columns["Date"] < len(values) else None
                if day_value is None:
                    continue
                if isinstance(day_value, (datetime, pd.Timestamp)):
                    parsed_date = pd.Timestamp(day_value)
                else:
                    numeric_day = pd.to_numeric(day_value, errors="coerce")
                    if pd.isna(numeric_day) or float(numeric_day) % 1:
                        continue
                    try:
                        parsed_date = pd.Timestamp(year=year, month=month_number, day=int(numeric_day))
                    except ValueError:
                        continue
                if parsed_date.strftime("%Y-%m") != month:
                    continue

                electricity = values[columns["Electricity"]] if columns["Electricity"] < len(values) else None
                water = values[columns["Water"]] if columns["Water"] < len(values) else None
                if electricity is None and water is None:
                    continue
                records.append({
                    "Date": parsed_date.strftime("%Y-%m-%d"),
                    "Occupancy": pd.NA,
                    "Room_Nights": pd.NA,
                    "Water": water,
                    "Electricity": electricity,
                    "Linen_Laundry": pd.NA,
                })
            if records:
                return _normalise_daily_frame(
                    pd.DataFrame(records),
                    month,
                    allow_missing_expenses=True,
                    allow_missing_room_nights=True,
                )
        raise ValueError("Could not find Tanggal/Date, Penggunaan Listrik/Electricity, and Penggunaan Air/Water columns.")
    finally:
        workbook.close()


def parse_daily_linen_cost_workbook(source, month: str) -> pd.DataFrame:
    """Extract daily linen cost, excluding zero-filled future rows beyond its data count."""
    month = _month_key(month)
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    try:
        for worksheet in workbook.worksheets:
            header_row = None
            date_col = None
            linen_total_col = None
            for row_idx, values in enumerate(worksheet.iter_rows(values_only=True), 1):
                labels = {
                    re.sub(r"[^a-z0-9]", "", str(value).strip().lower()): col_idx
                    for col_idx, value in enumerate(values)
                    if value is not None
                }
                possible_date = labels.get("tanggal")
                if possible_date is None:
                    possible_date = labels.get("date")
                total_col = next(
                    (idx for key, idx in labels.items() if "totalbiayaharian" in key),
                    None,
                )
                if possible_date is not None and total_col is not None:
                    header_row = row_idx
                    date_col = possible_date
                    linen_total_col = total_col
                    break
            if header_row is None:
                continue

            populated_days = None
            for values in worksheet.iter_rows(values_only=True):
                labels = [re.sub(r"[^a-z0-9]", "", str(value).strip().lower()) if value is not None else "" for value in values]
                if "jumlahhariyangadadata" in labels:
                    label_col = labels.index("jumlahhariyangadadata")
                    count_values = [
                        pd.to_numeric(value, errors="coerce")
                        for value in values[label_col + 1:]
                        if value is not None
                    ]
                    count = next((value for value in count_values if pd.notna(value)), None)
                    if count is not None:
                        populated_days = int(count)
                    break

            year, month_number = (int(part) for part in month.split("-"))
            records = []
            for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                day_value = values[date_col] if date_col < len(values) else None
                if day_value is None or str(day_value).strip().upper() == "TOTAL":
                    continue
                if isinstance(day_value, (datetime, pd.Timestamp)):
                    parsed_date = pd.Timestamp(day_value)
                else:
                    numeric_day = pd.to_numeric(day_value, errors="coerce")
                    if pd.isna(numeric_day) or float(numeric_day) % 1:
                        continue
                    try:
                        parsed_date = pd.Timestamp(year=year, month=month_number, day=int(numeric_day))
                    except ValueError:
                        continue
                if parsed_date.strftime("%Y-%m") != month:
                    continue
                if populated_days is not None and parsed_date.day > populated_days:
                    continue

                linen_total = values[linen_total_col] if linen_total_col < len(values) else None
                if linen_total is None:
                    continue
                records.append({
                    "Date": parsed_date.strftime("%Y-%m-%d"),
                    "Occupancy": pd.NA,
                    "Room_Nights": pd.NA,
                    "Water": pd.NA,
                    "Electricity": pd.NA,
                    "Linen_Laundry": linen_total,
                })
            if records:
                return _normalise_daily_frame(
                    pd.DataFrame(records),
                    month,
                    allow_missing_expenses=True,
                    allow_missing_room_nights=True,
                )
        raise ValueError("Could not find Tanggal/Date and TOTAL BIAYA HARIAN columns in the workbook.")
    finally:
        workbook.close()


def get_room_cost_storage_dir(storage_dir: Optional[str] = None) -> str:
    """Return the app-owned directory for persisted room-cost inputs."""
    path = storage_dir or os.environ.get("HOTEL_ROOM_COST_DIR")
    if not path:
        path = os.path.join(Path.home(), ".hotel-expense-tracker", "room_cost")
    return os.path.abspath(os.path.expanduser(path))


def get_room_cost_database_path(storage_dir: Optional[str] = None) -> str:
    """Return the SQLite database path used for operational data storage."""
    storage_path = get_room_cost_storage_dir(storage_dir)
    os.makedirs(storage_path, exist_ok=True)
    return os.path.join(storage_path, "hotel_data.sqlite3")


def _connect_room_cost_database(storage_dir: Optional[str] = None) -> sqlite3.Connection:
    connection = sqlite3.connect(get_room_cost_database_path(storage_dir))
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("""
        CREATE TABLE IF NOT EXISTS daily_room_costs (
            date TEXT PRIMARY KEY,
            month TEXT NOT NULL,
            occupancy REAL,
            room_nights INTEGER,
            water REAL,
            electricity REAL,
            linen_laundry REAL
        )
    """)
    connection.execute("""
        CREATE TABLE IF NOT EXISTS imported_files (
            sha256 TEXT PRIMARY KEY,
            source_path TEXT NOT NULL,
            source_name TEXT NOT NULL,
            imported_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            record_count INTEGER NOT NULL DEFAULT 0
        )
    """)
    connection.execute("""
        CREATE TABLE IF NOT EXISTS imported_rows (
            sha256 TEXT NOT NULL,
            sheet_name TEXT NOT NULL,
            row_number INTEGER NOT NULL,
            row_data TEXT NOT NULL,
            PRIMARY KEY (sha256, sheet_name, row_number),
            FOREIGN KEY (sha256) REFERENCES imported_files(sha256) ON DELETE CASCADE
        )
    """)
    connection.execute("""
        CREATE TABLE IF NOT EXISTS briefing_notes (
            date TEXT PRIMARY KEY,
            staff_incharge TEXT NOT NULL DEFAULT '',
            remarks TEXT NOT NULL DEFAULT ''
        )
    """)
    connection.execute("""
        CREATE TABLE IF NOT EXISTS manual_daily_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT NOT NULL,
            entry_type TEXT NOT NULL,
            staff_name TEXT NOT NULL DEFAULT '',
            section TEXT NOT NULL DEFAULT '',
            room_credit REAL,
            area TEXT NOT NULL DEFAULT '',
            description TEXT NOT NULL DEFAULT '',
            qty REAL,
            unit TEXT NOT NULL DEFAULT '',
            unit_cost REAL,
            total REAL,
            status TEXT NOT NULL DEFAULT '',
            pic TEXT NOT NULL DEFAULT '',
            notes TEXT NOT NULL DEFAULT ''
        )
    """)
    connection.commit()
    return connection


def _upsert_room_cost_rows(
    connection: sqlite3.Connection,
    data: pd.DataFrame,
    *,
    ignore_existing: bool = False,
) -> int:
    if data.empty:
        return 0
    command = "INSERT OR IGNORE" if ignore_existing else "INSERT OR REPLACE"
    sql = f"""{command} INTO daily_room_costs
        (date, month, occupancy, room_nights, water, electricity, linen_laundry)
        VALUES (?, ?, ?, ?, ?, ?, ?)"""
    records = []
    for row in data.to_dict(orient="records"):
        records.append((
            row["Date"], row["Date"][:7],
            None if pd.isna(row["Occupancy"]) else float(row["Occupancy"]),
            None if pd.isna(row["Room_Nights"]) else int(row["Room_Nights"]),
            None if pd.isna(row["Water"]) else float(row["Water"]),
            None if pd.isna(row["Electricity"]) else float(row["Electricity"]),
            None if pd.isna(row["Linen_Laundry"]) else float(row["Linen_Laundry"]),
        ))
    before = connection.total_changes
    connection.executemany(sql, records)
    return connection.total_changes - before


def _legacy_daily_room_costs(month: str, storage_dir: Optional[str] = None) -> pd.DataFrame:
    path = os.path.join(get_room_cost_storage_dir(storage_dir), f"{_month_key(month)}.json")
    if not os.path.exists(path):
        return pd.DataFrame(columns=DAILY_COLUMNS)
    try:
        with open(path, "r", encoding="utf-8") as file_obj:
            rows = json.load(file_obj)
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Could not read saved room-cost data for {month}.") from exc
    return _normalise_daily_frame(
        pd.DataFrame(rows), month, allow_missing_expenses=True, allow_missing_room_nights=True
    )


def _import_workbook_rows(connection: sqlite3.Connection, source_path: str) -> dict:
    path = os.path.abspath(os.path.expanduser(source_path))
    digest = hashlib.sha256()
    with open(path, "rb") as source_file:
        for chunk in iter(lambda: source_file.read(1024 * 1024), b""):
            digest.update(chunk)
    file_hash = digest.hexdigest()
    if connection.execute("SELECT 1 FROM imported_files WHERE sha256 = ?", (file_hash,)).fetchone():
        return {"file": path, "imported": False, "rows": 0}

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    records = []
    try:
        for worksheet in workbook.worksheets:
            for row_number, values in enumerate(worksheet.iter_rows(values_only=True), 1):
                if not any(value is not None for value in values):
                    continue
                serializable = []
                for value in values:
                    if value is None or (not isinstance(value, (str, bytes)) and pd.isna(value)):
                        serializable.append(None)
                    elif isinstance(value, (datetime, pd.Timestamp)):
                        serializable.append(pd.Timestamp(value).isoformat())
                    elif hasattr(value, "item"):
                        serializable.append(value.item())
                    else:
                        serializable.append(value)
                records.append((file_hash, worksheet.title, row_number, json.dumps(serializable, ensure_ascii=True)))
    finally:
        workbook.close()

    connection.execute(
        "INSERT INTO imported_files (sha256, source_path, source_name, record_count) VALUES (?, ?, ?, ?)",
        (file_hash, path, os.path.basename(path), len(records)),
    )
    connection.executemany(
        "INSERT INTO imported_rows (sha256, sheet_name, row_number, row_data) VALUES (?, ?, ?, ?)",
        records,
    )
    return {"file": path, "imported": True, "rows": len(records)}


def migrate_existing_data_to_database(
    month: str,
    source_files: Optional[list[str]] = None,
    storage_dir: Optional[str] = None,
) -> dict:
    """Copy existing daily JSON and selected workbook rows into the local SQLite store."""
    month = _month_key(month)
    legacy_data = _legacy_daily_room_costs(month, storage_dir)
    connection = _connect_room_cost_database(storage_dir)
    try:
        room_cost_rows = _upsert_room_cost_rows(connection, legacy_data, ignore_existing=True)
        workbook_results = []
        for source_path in source_files or []:
            if not os.path.isfile(os.path.expanduser(source_path)):
                raise FileNotFoundError(f"Source workbook not found: {source_path}")
            workbook_results.append(_import_workbook_rows(connection, source_path))
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    return {
        "database": get_room_cost_database_path(storage_dir),
        "room_cost_rows_imported": room_cost_rows,
        "workbooks": workbook_results,
    }


def load_manual_briefing_data(report_date, storage_dir: Optional[str] = None) -> dict:
    """Load manual briefing fields and related rows for one date."""
    date_value = pd.Timestamp(report_date).strftime("%Y-%m-%d")
    connection = _connect_room_cost_database(storage_dir)
    try:
        note = connection.execute(
            "SELECT staff_incharge, remarks FROM briefing_notes WHERE date = ?", (date_value,)
        ).fetchone()
        entries = connection.execute(
            "SELECT entry_type, staff_name, section, room_credit, area, description, qty, unit, "
            "unit_cost, total, status, pic, notes FROM manual_daily_entries WHERE date = ? ORDER BY id",
            (date_value,),
        ).fetchall()
    finally:
        connection.close()

    entry_columns = [
        "Entry_Type", "Staff_Name", "Section", "Room_Credit", "Area", "Description",
        "Qty", "Unit", "Unit_Cost", "Total", "Status", "PIC", "Notes",
    ]
    entries_frame = pd.DataFrame(entries, columns=entry_columns)
    return {
        "Staff_Incharge": note[0] if note else "",
        "Remarks": note[1] if note else "",
        "Entries": entries_frame,
    }


def save_manual_briefing_data(
    report_date,
    *,
    staff_incharge: str = "",
    remarks: str = "",
    entries: Optional[pd.DataFrame] = None,
    storage_dir: Optional[str] = None,
) -> None:
    """Replace manual briefing entries for a date in one SQLite transaction."""
    date_value = pd.Timestamp(report_date).strftime("%Y-%m-%d")
    allowed_types = {"ra", "project_pa", "daily_pa", "guest_laundry"}
    records = []
    if entries is not None and not entries.empty:
        for row in entries.to_dict(orient="records"):
            entry_type = str(row.get("Entry_Type", "")).strip()
            if entry_type not in allowed_types:
                raise ValueError(f"Unsupported manual entry type: {entry_type}")
            text = lambda key: "" if pd.isna(row.get(key)) else str(row.get(key)).strip()
            number = lambda key: None if pd.isna(row.get(key)) or row.get(key) == "" else float(row[key])
            qty = number("Qty")
            unit_cost = number("Unit_Cost")
            total = number("Total")
            if total is None and qty is not None and unit_cost is not None:
                total = qty * unit_cost
            records.append((
                date_value, entry_type, text("Staff_Name"), text("Section"),
                number("Room_Credit"), text("Area"), text("Description"), qty,
                text("Unit"), unit_cost, total, text("Status"), text("PIC"), text("Notes"),
            ))

    connection = _connect_room_cost_database(storage_dir)
    try:
        with connection:
            connection.execute(
                "INSERT INTO briefing_notes (date, staff_incharge, remarks) VALUES (?, ?, ?) "
                "ON CONFLICT(date) DO UPDATE SET staff_incharge = excluded.staff_incharge, remarks = excluded.remarks",
                (date_value, str(staff_incharge or "").strip(), str(remarks or "").strip()),
            )
            connection.execute("DELETE FROM manual_daily_entries WHERE date = ?", (date_value,))
            connection.executemany(
                "INSERT INTO manual_daily_entries "
                "(date, entry_type, staff_name, section, room_credit, area, description, qty, unit, unit_cost, total, status, pic, notes) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                records,
            )
    finally:
        connection.close()


def load_daily_room_costs(month: str, storage_dir: Optional[str] = None) -> pd.DataFrame:
    """Load daily costs from SQLite, importing legacy month JSON if needed."""
    month = _month_key(month)
    legacy_data = _legacy_daily_room_costs(month, storage_dir)
    connection = _connect_room_cost_database(storage_dir)
    try:
        _upsert_room_cost_rows(connection, legacy_data, ignore_existing=True)
        rows = connection.execute(
            "SELECT date, occupancy, room_nights, water, electricity, linen_laundry "
            "FROM daily_room_costs WHERE month = ? ORDER BY date",
            (month,),
        ).fetchall()
        connection.commit()
    finally:
        connection.close()
    frame = pd.DataFrame(rows, columns=DAILY_COLUMNS)
    return _normalise_daily_frame(
        frame, month, allow_missing_expenses=True, allow_missing_room_nights=True
    )


def save_daily_room_costs(month: str, daily_df: pd.DataFrame, storage_dir: Optional[str] = None) -> None:
    """Persist validated daily costs in SQLite, retaining legacy JSON as backup."""
    month = _month_key(month)
    data = _normalise_daily_frame(
        daily_df, month, allow_missing_expenses=True, allow_missing_room_nights=True
    )
    connection = _connect_room_cost_database(storage_dir)
    try:
        _upsert_room_cost_rows(connection, data)
        connection.commit()
    finally:
        connection.close()


def calculate_mtd_occupancy_and_room_nights(
    daily_df: pd.DataFrame,
    month: str,
) -> tuple[float, int]:
    """Return average available occupancy and cumulative room nights for a month."""
    daily = _normalise_daily_frame(
        daily_df, month, allow_missing_expenses=True, allow_missing_room_nights=True
    )
    occupancy = pd.to_numeric(daily["Occupancy"], errors="coerce")
    room_nights = pd.to_numeric(daily["Room_Nights"], errors="coerce")
    occupancy_mtd = float(occupancy.mean()) if occupancy.notna().any() else float("nan")
    room_nights_mtd = int(room_nights.sum()) if room_nights.notna().any() else 0
    return occupancy_mtd, room_nights_mtd


def merge_daily_room_costs(
    saved_df: pd.DataFrame,
    imported_df: pd.DataFrame,
    replace_dates: bool = False,
    month: Optional[str] = None,
) -> pd.DataFrame:
    """Merge imported dates into saved rows, optionally replacing duplicates."""
    if month is None:
        if imported_df.empty and saved_df.empty:
            raise ValueError("Month is required when both datasets are empty.")
        date_values = pd.concat([saved_df.get("Date", pd.Series(dtype=str)), imported_df.get("Date", pd.Series(dtype=str))])
        month = str(date_values.iloc[0])[:7]
    month = _month_key(month)
    saved = _normalise_daily_frame(
        saved_df, month, allow_missing_expenses=True, allow_missing_room_nights=True
    )
    imported = _normalise_daily_frame(
        imported_df, month, allow_missing_expenses=True, allow_missing_room_nights=True
    )
    duplicates = set(saved["Date"]) & set(imported["Date"])
    if duplicates and not replace_dates:
        dates = ", ".join(sorted(duplicates))
        raise ValueError(f"Date(s) already saved; confirm replacement before importing: {dates}.")
    preserved = saved[saved["Date"].isin(duplicates)].set_index("Date")
    replacements = imported.set_index("Date")
    for column in DAILY_COLUMNS[1:]:
        if column in preserved.columns:
            blank_values = replacements[column].isna()
            shared_dates = replacements.index.intersection(preserved.index)
            replacements.loc[shared_dates[blank_values.loc[shared_dates]], column] = preserved.loc[
                shared_dates[blank_values.loc[shared_dates]], column
            ]
    replacements = replacements.reset_index()
    combined = pd.concat(
        [saved[~saved["Date"].isin(duplicates)], replacements], ignore_index=True
    )
    return combined.sort_values("Date").reset_index(drop=True)


def calculate_daily_room_costs(
    sr_df: pd.DataFrame,
    daily_df: pd.DataFrame,
    month: str,
) -> tuple[pd.DataFrame, pd.DataFrame, list[str]]:
    """Calculate per-day costs and the current retroactive MTD SR rate."""
    month = _month_key(month)
    daily = _normalise_daily_frame(
        daily_df, month, allow_missing_expenses=True, allow_missing_room_nights=True
    )
    if daily.empty:
        empty = pd.DataFrame()
        return empty, empty, []

    warnings = []
    missing_expenses = daily[EXPENSE_COLUMNS].isna()
    for column in EXPENSE_COLUMNS:
        missing_count = int(missing_expenses[column].sum())
        if missing_count:
            warnings.append(
                f"{column.replace('_', ' ')} is missing for {missing_count} date(s); "
                "combined room cost will remain incomplete until those amounts are entered."
            )
    missing_room_nights = int(daily["Room_Nights"].isna().sum())
    if missing_room_nights:
        warnings.append(
            f"Room sold is missing for {missing_room_nights} date(s); per-room calculations remain incomplete for those dates."
        )
    report_end = pd.Timestamp(daily["Date"].max())
    expected_dates = set(pd.date_range(f"{month}-01", report_end, freq="D").strftime("%Y-%m-%d"))
    missing_dates = sorted(expected_dates - set(daily["Date"]))
    if missing_dates:
        warnings.append(
            f"Room-night input is missing {len(missing_dates)} date(s) through {report_end:%Y-%m-%d}; "
            "the MTD SR rate uses only saved dates."
        )
    sr_by_date = pd.Series(dtype=float)
    if sr_df is not None and not sr_df.empty:
        if "Date" not in sr_df.columns or "Total" not in sr_df.columns:
            raise ValueError("SR data must contain Date and Total columns.")
        sr_dates = pd.to_datetime(sr_df["Date"], errors="coerce", format="mixed")
        sr_totals = pd.to_numeric(sr_df["Total"], errors="coerce")
        invalid_sr = sr_dates.isna() | sr_totals.isna()
        if invalid_sr.any():
            warnings.append(f"Ignored {int(invalid_sr.sum())} SR row(s) with missing or invalid dates/amounts.")
        valid = ~invalid_sr
        outside_month = valid & (sr_dates.dt.strftime("%Y-%m") != month)
        if outside_month.any():
            warnings.append(f"Ignored {int(outside_month.sum())} SR row(s) outside {month}.")
        valid &= ~outside_month
        after_report_end = valid & (sr_dates.dt.normalize() > report_end)
        if after_report_end.any():
            warnings.append(f"Ignored {int(after_report_end.sum())} SR row(s) after the latest daily input date ({report_end:%Y-%m-%d}).")
        valid &= ~after_report_end
        if (sr_totals[valid] < 0).any():
            raise ValueError("SR totals cannot be negative.")
        if valid.any():
            sr_by_date = sr_totals[valid].groupby(sr_dates[valid].dt.strftime("%Y-%m-%d")).sum()

    result = daily.copy()
    result["SR_Received_On_Date"] = result["Date"].map(sr_by_date).fillna(0.0)
    mtd_sr_spend = float(sr_by_date.sum())
    mtd_room_nights = int(result["Room_Nights"].sum())
    sr_rate = mtd_sr_spend / mtd_room_nights

    for column in EXPENSE_COLUMNS:
        result[f"{column}_Per_Room"] = result[column] / result["Room_Nights"]
    result["MTD_SR_Spend"] = mtd_sr_spend
    result["MTD_Room_Nights"] = mtd_room_nights
    result["SR_Per_Room"] = sr_rate
    sr_category_columns = {
        "Guest Supplies": "Guest_Supplies_Per_Room",
        "Cleaning Supplies": "Cleaning_Supplies_Per_Room",
        "Paper Supplies": "Paper_Supplies_Per_Room",
        "Print & Stationery": "Print_Stationery_Per_Room",
        "Other": "Other_SR_Per_Room",
    }
    category_totals = {category: 0.0 for category in sr_category_columns}
    if sr_df is not None and not sr_df.empty and "Category" in sr_df.columns and valid.any():
        normalized_categories = sr_df.loc[valid, "Category"].fillna("").astype(str).str.strip().str.lower()
        category_aliases = {
            "guest supplies": "Guest Supplies",
            "cleaning supplies": "Cleaning Supplies",
            "paper supplies": "Paper Supplies",
            "print": "Print & Stationery",
            "print & stationery": "Print & Stationery",
        }
        categorized = pd.DataFrame({
            "Category": normalized_categories.map(category_aliases).fillna("Other"),
            "Total": sr_totals.loc[valid],
        })
        category_totals.update(categorized.groupby("Category")["Total"].sum().to_dict())
    else:
        category_totals["Other"] = mtd_sr_spend
    for category, column in sr_category_columns.items():
        result[column] = category_totals[category] / mtd_room_nights
    expense_complete = result[EXPENSE_COLUMNS].notna().all(axis=1)
    room_cost_complete = expense_complete & result["Room_Nights"].notna()
    result["Total_Cost_Per_Room"] = (
        result[[*sr_category_columns.values(), *(f"{column}_Per_Room" for column in EXPENSE_COLUMNS)]].sum(
            axis=1, min_count=len(sr_category_columns) + len(EXPENSE_COLUMNS)
        )
    )
    result.loc[~room_cost_complete, "Total_Cost_Per_Room"] = float("nan")
    result["Total_Room_Cost"] = (
        result[EXPENSE_COLUMNS].sum(axis=1, min_count=len(EXPENSE_COLUMNS))
        + result["SR_Per_Room"] * result["Room_Nights"]
    )
    result.loc[~room_cost_complete, "Total_Room_Cost"] = float("nan")

    monthly_total_room_cost = (
        float(result["Total_Room_Cost"].sum()) if room_cost_complete.all() else float("nan")
    )

    monthly = pd.DataFrame([{
        "Month": month,
        "Room_Nights": mtd_room_nights,
        "Occupancy_Average": float(result["Occupancy"].mean()) if result["Occupancy"].notna().any() else float("nan"),
        "Water_Total": float(result["Water"].sum(min_count=1)) if result["Water"].notna().any() else float("nan"),
        "Electricity_Total": float(result["Electricity"].sum(min_count=1)) if result["Electricity"].notna().any() else float("nan"),
        "Linen_Laundry_Total": float(result["Linen_Laundry"].sum(min_count=1)) if result["Linen_Laundry"].notna().any() else float("nan"),
        "SR_Total": mtd_sr_spend,
        "Total_Room_Cost": monthly_total_room_cost,
        "Water_Per_Room": float(result.loc[result["Room_Nights"].notna(), "Water"].sum(min_count=1)) / mtd_room_nights if result["Room_Nights"].notna().all() and result["Water"].notna().all() else float("nan"),
        "Electricity_Per_Room": float(result.loc[result["Room_Nights"].notna(), "Electricity"].sum(min_count=1)) / mtd_room_nights if result["Room_Nights"].notna().all() and result["Electricity"].notna().all() else float("nan"),
        "Linen_Laundry_Per_Room": float(result.loc[result["Room_Nights"].notna(), "Linen_Laundry"].sum(min_count=1)) / mtd_room_nights if result["Room_Nights"].notna().all() and result["Linen_Laundry"].notna().all() else float("nan"),
        "SR_Per_Room": sr_rate,
        "Total_Cost_Per_Room": monthly_total_room_cost / mtd_room_nights,
    }])
    return result, monthly, warnings